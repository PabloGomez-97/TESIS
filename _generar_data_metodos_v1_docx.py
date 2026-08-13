"""Genera Data y Métodos v1 a partir del outline y el estilo de Introducción v1."""
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.shared import Pt, Cm
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

OUT_DOWNLOADS = Path("/Users/pablogomez/Downloads/Data_y_Metodos_v1.docx")
OUT_REPO = Path(__file__).resolve().parent / "Data_y_Metodos_v1.docx"
INTRO = Path("/Users/pablogomez/Downloads/Introducción v1.docx")


def set_run_font(run, size=11, bold=False, italic=False):
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic


def add_heading_custom(doc, text, level=1):
    h = doc.add_heading(text, level=level)
    for run in h.runs:
        set_run_font(run, size=14 if level == 1 else 12, bold=True)
    return h


def add_para(doc, text, *, first_line_indent=True):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_after = Pt(8)
    pf.space_before = Pt(0)
    pf.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    if first_line_indent:
        pf.first_line_indent = Cm(1.25)
    run = p.add_run(text)
    set_run_font(run, size=11)
    return p


def add_note(doc, text):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_after = Pt(10)
    pf.space_before = Pt(4)
    run = p.add_run(text)
    set_run_font(run, size=10, italic=True)
    return p


def build_methods_doc():
    doc = Document()

    section = doc.sections[0]
    section.top_margin = Cm(2.5)
    section.bottom_margin = Cm(2.5)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2.5)

    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(11)

    add_heading_custom(doc, "2. Datos y métodos", level=1)

    # --- 2.1 ---
    add_heading_custom(doc, "2.1 Conjuntos de datos, dominio y períodos", level=2)

    add_para(
        doc,
        "El análisis utiliza dos fuentes de precipitación diaria. Como referencia "
        "observacional del período histórico se emplea CR2MET v2.5 (Centro de Ciencia "
        "del Clima y la Resiliencia, CR2), un producto gridded de precipitación diaria "
        "sobre Chile en una grilla regular de latitud–longitud (~0,05°), con unidades "
        "en mm día⁻¹. Como modelo climático regional se utiliza ALADIN CHP12 "
        "(configuración CNRM-ALADIN64), cuya precipitación se entrega en kg m⁻² s⁻¹ y "
        "se convierte a mm día⁻¹ multiplicando por 86 400. El dominio de estudio es "
        "Chile continental; la máscara terrestre se construye a partir de la geometría "
        "de Chile en Natural Earth (10 m). Sobre la grilla nativa de ALADIN, el dominio "
        "continental comprende aproximadamente 3 587 celdas válidas.",
    )

    add_para(
        doc,
        "Para permitir comparaciones píxel a píxel, CR2MET se interpola linealmente a los "
        "centros de celda de la grilla nativa de ALADIN. Esta decisión fija ALADIN como "
        "marco espacial de referencia, porque es el producto que se evalúa en el histórico "
        "y el único disponible para el futuro. La interpolación lineal es un método "
        "estándar y reproducible; se reconoce, no obstante, que introduce un suavizado "
        "leve de CR2MET y que no conserva exactamente la precipitación integrada del "
        "dominio.",
    )

    add_para(
        doc,
        "Se analizan dos ventanas de 35 años de igual longitud: el histórico 1980–2014 "
        "(CR2MET y ALADIN) y el futuro 2040–2074 bajo el escenario SSP5-8.5 (solo ALADIN). "
        "El emparejamiento de longitudes evita que diferencias en el tamaño muestral "
        "condicionen la comparación histórico–futuro. Los análisis regionales de "
        "distribuciones de duración se concentran en cuatro cuencas hidrográficas "
        "representativas del gradiente climático chileno —Loa (norte árido), Maipo y "
        "Maule (zona central) y Biobío (sur más húmedo)—, delimitadas como ventanas "
        "rectangulares de longitud–latitud sobre la grilla de ALADIN e intersectadas con "
        "la máscara de Chile continental. Los risk ratios, en cambio, se calculan sobre "
        "pools regionales definidos por regiones administrativas (Coquimbo, O’Higgins, "
        "La Araucanía y Los Lagos), usando geometrías Natural Earth admin_1.",
    )

    add_note(
        doc,
        "→ Tabla 1: resumen de datasets, grilla, períodos, variables y umbrales. "
        "→ Fig. 2: mapa del dominio con cuencas (PDFs) y regiones administrativas (RR).",
    )

    # --- 2.2 ---
    add_heading_custom(doc, "2.2 Definición de días húmedos/secos y dry spells", level=2)

    add_para(
        doc,
        "Se define wet day (día húmedo) como un día con precipitación diaria "
        "pr ≥ R, donde R es un umbral en mm día⁻¹, y dry day (día seco) como un día con "
        "pr < R. Un dry spell es una racha consecutiva de dry days; su duración t es el "
        "número de días de la racha. Los días con valor faltante (NaN) no se tratan como "
        "secos. Las rachas que comienzan en el primer día o terminan en el último día de "
        "la serie temporal de cada período se excluyen del análisis, para evitar spells "
        "censurados por el borde de la ventana.",
    )

    add_para(
        doc,
        "La comparación histórica entre CR2MET y ALADIN no puede realizarse de forma "
        "equitativa con un mismo umbral fijo. Con R = 1 mm día⁻¹ en ambos productos, "
        "ALADIN presenta una fracción espacial media de días húmedos sustancialmente "
        "mayor que CR2MET, lo que fragmenta las rachas secas y genera spells "
        "artificialmente más cortos y numerosos. Por ello, el análisis principal adopta "
        "el criterio i (umbral global calibrado): CR2MET con R = 1 mm día⁻¹ y ALADIN con "
        "un umbral calibrado R₀ = 5,285 mm día⁻¹, elegido para igualar la fracción "
        "espacial media de wet days de CR2MET a 1 mm (véase §2.3). Denotamos este umbral "
        "como R₀ —y no como τ*— para evitar confusión con el exponente τ_P de leyes de "
        "potencia en PDFs de precipitación diaria (Martinez-Villalobos et al., 2022).",
    )

    add_para(
        doc,
        "Como análisis de sensibilidad se consideran además el criterio ii (mismo umbral: "
        "1 mm día⁻¹ en CR2MET y ALADIN) y el criterio iii (calibración local: R₀(x,y) por "
        "píxel, de modo que en cada celda la fracción temporal de wet days de ALADIN "
        "iguale la de CR2MET a 1 mm). El criterio i es el de referencia para resultados "
        "principales; los criterios ii y iii permiten separar cuánto del sesgo histórico "
        "se debe a la definición de día húmedo y cuánto permanece como diferencia "
        "estructural del modelo.",
    )

    # --- 2.3 ---
    add_heading_custom(doc, "2.3 Calibración del umbral R₀", level=2)

    add_para(
        doc,
        "La calibración de R₀ busca el valor tal que el promedio espacial sobre Chile "
        "continental de la fracción temporal de días con pr ≥ R en ALADIN iguale la "
        "fracción correspondiente de CR2MET con R = 1 mm día⁻¹. Formalmente, si "
        "F(R) denota la fracción espacial media de wet days, se resuelve "
        "F_ALADIN(R₀) = F_CR2MET(1 mm). En el período 1980–2014, CR2MET a 1 mm produce "
        "F = 20,94%, mientras que ALADIN a 1 mm produce F = 30,68% (+9,7 puntos "
        "porcentuales). El umbral R₀ = 5,285 mm día⁻¹ reproduce F = 20,94% en ALADIN, "
        "con error inferior a 0,0001 puntos porcentuales.",
    )

    add_para(
        doc,
        "Esta calibración sigue la lógica de emparejamiento de días húmedos empleada en "
        "estudios distributivos de precipitación (Martinez-Villalobos et al., 2022): "
        "igualar la frecuencia de wet days antes de comparar estadísticas de rachas o "
        "intensidades. R₀ es, por tanto, un paso metodológico previo y no la pregunta "
        "científica central del trabajo. Cabe subrayar qué no garantiza la calibración: "
        "igualar la fracción de wet days no implica igualdad en intensidad de "
        "precipitación, en variabilidad interanual ni en la distribución de duraciones "
        "de dry spells. De hecho, incluso con R₀, ALADIN sigue subestimando la "
        "persistencia histórica respecto a CR2MET (véase §3), lo que indica un sesgo "
        "residual del modelo más allá del umbral.",
    )

    add_para(
        doc,
        "La variante local R₀(x,y) (criterio iii) calcula, en cada píxel, el umbral de "
        "ALADIN que iguala la fracción temporal de wet days de CR2MET a 1 mm en ese "
        "mismo píxel. Esta variante reduce residuos espaciales del sesgo de wet days, "
        "pero introduce un campo de umbrales heterogéneo y mayor complejidad "
        "interpretativa; por ello se reserva a sensibilidad y material suplementario.",
    )

    add_note(
        doc,
        "→ Fig. 3 / fig_p6_ftau: curva F(R) de ALADIN, punto CR2MET @ 1 mm y R₀ = 5,285 mm. "
        "→ Fig. umbral 1 mm vs R₀: efecto sobre wet days, duración media y PDFs. "
        "→ Fig. S2: mapa de R₀ local.",
    )

    # --- 2.4 ---
    add_heading_custom(doc, "2.4 Construcción y normalización de las PDFs", level=2)

    add_para(
        doc,
        "Para caracterizar la distribución de duraciones se construyen PDFs empíricas "
        "de la variable t (días). Los bins se definen aproximadamente uniformes en el "
        "espacio log(t), de modo que la representación natural de las curvas sea en "
        "escala log–log, siguiendo la convención usada en PDFs de precipitación diaria "
        "(Martinez-Villalobos & Neelin, 2018; Martinez-Villalobos et al., 2022). En cada "
        "bin, la altura f(t) se obtiene como el cociente entre el número de spells "
        "contados en el intervalo y el ancho del bin.",
    )

    add_para(
        doc,
        "A diferencia de una densidad de probabilidad clásica —cuya integral respecto "
        "de t iguala 1—, aquí se normaliza de modo que",
    )

    # Equation-like paragraph centered
    p_eq = doc.add_paragraph()
    p_eq.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_eq.paragraph_format.space_before = Pt(6)
    p_eq.paragraph_format.space_after = Pt(6)
    run = p_eq.add_run("(1)\t∫₀^∞ f(t) dt = n_spells")
    set_run_font(run, size=11, italic=True)

    add_para(
        doc,
        "donde n_spells es el número total de dry spells del conjunto considerado "
        "(cuenca × período × producto). Esta convención preserva simultáneamente la "
        "forma de la distribución y el número total de eventos. Normalizar a 1 "
        "ocultaría si el futuro produce más o menos spells que el histórico, "
        "información relevante cuando cambian a la vez la frecuencia de eventos y la "
        "cola de duraciones largas. La PDF de duraciones no debe confundirse con una "
        "tendencia temporal: la primera describe cómo se reparten las rachas en un "
        "pool de eventos; la segunda describe la evolución año a año de un estadístico.",
    )

    add_para(
        doc,
        "Como complemento escalar a las PDFs se reportan, para cada cuenca y período, "
        "la media, la mediana, los percentiles 90 y 99 (p90, p99) y el máximo del pool "
        "regional de duraciones.",
    )

    # --- 2.5 ---
    add_heading_custom(doc, "2.5 Agrupación regional y estadísticas espaciales", level=2)

    add_para(
        doc,
        "Para las PDFs por cuenca se adopta un pool regional: todos los dry spells de "
        "todos los píxeles situados dentro de la ventana de la cuenca se agrupan en un "
        "único conjunto de eventos. Este enfoque, en la línea de análisis regionales "
        "de Martinez-Villalobos & Neelin (2018), aumenta el tamaño muestral y permite "
        "resolver mejor la forma de la distribución —en particular la cola—, a costa "
        "de mezclar heterogeneidad espacial dentro de la cuenca. Las cuencas se "
        "seleccionaron como casos de estudio a lo largo del gradiente norte–sur, no "
        "como delineación hidrológica oficial; la delimitación mediante cajas "
        "lon/lat sobre la grilla ALADIN es explícita y reproducible.",
    )

    add_para(
        doc,
        "Para los mapas espaciales se calculan, en cada píxel, la duración media de "
        "los dry spells y el percentil 99 de las duraciones (t99), y se reportan las "
        "diferencias Δ = ALADIN − CR2MET en el período histórico 1980–2014. La "
        "significancia de estas diferencias se evalúa mediante bootstrap por años: "
        "se remuestrean años completos (no días sueltos) para respetar la "
        "autocorrelación temporal, se reconstruye el pool de spells de los años "
        "sorteados y se recalcula la métrica en cada réplica (1 000 réplicas). Una "
        "diferencia se considera significativa al 95% cuando el intervalo de "
        "confianza de Δ no incluye el cero.",
    )

    # --- 2.6 ---
    add_heading_custom(doc, "2.6 Risk ratios y cambios futuros", level=2)

    add_para(
        doc,
        "El cambio proyectado en la probabilidad de dry spells largos se cuantifica "
        "con el risk ratio",
    )

    p_eq2 = doc.add_paragraph()
    p_eq2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_eq2.paragraph_format.space_before = Pt(6)
    p_eq2.paragraph_format.space_after = Pt(6)
    run = p_eq2.add_run("(2)\tRR(D) = P_fut(t ≥ D) / P_hist(t ≥ D)")
    set_run_font(run, size=11, italic=True)

    add_para(
        doc,
        "donde P(t ≥ D) es la fracción de dry spells cuya duración es al menos D días, "
        "evaluada en ALADIN futuro (2040–2074) y ALADIN histórico (1980–2014) con el "
        "mismo umbral R₀ = 5,285 mm día⁻¹. Un RR(D) > 1 indica mayor probabilidad de "
        "spells largos en el futuro; un RR(D) < 1, menor probabilidad. Como el futuro "
        "no existe en CR2MET, esta comparación se realiza necesariamente dentro del "
        "mismo modelo (ALADIN hist versus ALADIN fut).",
    )

    add_para(
        doc,
        "Los risk ratios se estiman sobre pools regionales por región administrativa, "
        "considerando spells cuyo día de inicio cae entre marzo y noviembre, de modo "
        "de enfocar la estación relevante para rachas que condicionan la temporada "
        "seca y de transición. Se calculan curvas RR versus D (umbrales de duración) "
        "e intervalos de confianza al 95% mediante bootstrap por años (1 000 "
        "réplicas) sobre el pool regional. Se interpreta como aumento significativo "
        "cuando RR > 1 y el IC95 queda completamente por encima de 1. Cuando la "
        "probabilidad histórica de spells ≥ D es muy baja —como ocurre en regiones "
        "húmedas del sur para D grandes—, el cociente RR puede volverse grande e "
        "inestable; el IC95 amplio refleja entonces incertidumbre por rareza de "
        "eventos, no necesariamente una señal climática robusta.",
    )

    add_note(
        doc,
        "→ Fig. 8 / rr_curves_aladin: RR(D) por región con IC95 (bootstrap por años).",
    )

    # --- 2.7 ---
    add_heading_custom(doc, "2.7 Sensibilidad a la elección del umbral", level=2)

    add_para(
        doc,
        "Para documentar que las conclusiones históricas dependen del umbral, se "
        "comparan en cuencas seleccionadas (Loa, Maule y Biobío) cuatro series en el "
        "período 1980–2014: (i) CR2MET a 1 mm, (ii) ALADIN a 1 mm, (iii) ALADIN a R₀ "
        "global (5,285 mm) y (iv) ALADIN a R₀ local. Sin calibración (ALADIN a 1 mm), "
        "el modelo produce muchos más spells y de menor duración —por ejemplo, en la "
        "cuenca del Loa la media pasa de ~66 días en CR2MET a ~27 días en ALADIN—, "
        "distorsionando la comparación. Con R₀ global o local, las PDFs de ALADIN se "
        "acercan sustancialmente a CR2MET tanto en forma como en número de eventos, "
        "lo que justifica adoptar el criterio i como configuración principal del "
        "análisis. El detalle gráfico de esta sensibilidad se presenta en material "
        "suplementario junto con el mapa de R₀ local.",
    )

    add_note(
        doc,
        "→ Fig. 4 / fig_umbral_1mm_vs_R0_*: comparación 1 mm versus R₀ (wet days, "
        "duración media y PDFs). → Fig. S3: calibración Maule y Biobío.",
    )

    # Closing bridge
    add_para(
        doc,
        "En síntesis, la cadena metodológica es: (1) alinear espacialmente CR2MET y "
        "ALADIN; (2) definir wet/dry days y dry spells; (3) calibrar R₀ para una "
        "comparación histórica equitativa; (4) construir PDFs con ∫ f dt = n_spells "
        "mediante pool regional por cuenca; (5) mapear media y t99 con significancia "
        "por bootstrap; y (6) cuantificar el cambio futuro con RR(D) e IC95. Las "
        "secciones siguientes aplican este marco a la evaluación histórica (§3) y a "
        "las proyecciones SSP5-8.5 (§4).",
    )

    return doc


def merge_into_intro(methods_doc):
    """Inserta la sección Data y Métodos en Introducción v1 (reemplaza heading vacío)."""
    if not INTRO.exists():
        return None
    intro = Document(str(INTRO))

    # Find "Data y métodos" heading index and remove empty content until Referencias
    paras = intro.paragraphs
    start_idx = None
    end_idx = None
    for i, p in enumerate(paras):
        t = p.text.strip().lower()
        if start_idx is None and ("data y método" in t or "datos y método" in t):
            start_idx = i
        if start_idx is not None and i > start_idx and t.startswith("referencia"):
            end_idx = i
            break

    if start_idx is None:
        return None

    # Build new document: copy intro until start, then methods, then refs
    out = Document()
    for s in out.sections:
        s.top_margin = Cm(2.5)
        s.bottom_margin = Cm(2.5)
        s.left_margin = Cm(2.5)
        s.right_margin = Cm(2.5)

    def copy_paragraph(src_p):
        # Simple text copy preserving heading level roughly
        style_name = src_p.style.name if src_p.style else "Normal"
        text = src_p.text
        if not text.strip() and "Heading" not in style_name:
            return
        if style_name.startswith("Heading"):
            level = int(style_name.replace("Heading ", "")) if style_name[-1].isdigit() else 1
            add_heading_custom(out, text, level=level)
        else:
            add_para(out, text, first_line_indent=True)

    for i, p in enumerate(paras):
        if i == start_idx:
            break
        copy_paragraph(p)

    # Append methods body (skip duplicate top heading "2. Datos..." — use "Datos y métodos")
    # We'll add methods content from methods_doc
    first_heading_skipped = False
    for p in methods_doc.paragraphs:
        text = p.text
        style_name = p.style.name if p.style else "Normal"
        if not first_heading_skipped and style_name.startswith("Heading") and "Datos y método" in text:
            add_heading_custom(out, "Datos y métodos", level=1)
            first_heading_skipped = True
            continue
        if style_name.startswith("Heading"):
            level = int(style_name.replace("Heading ", "")) if style_name[-1].isdigit() else 2
            # methods uses 2.x as level 2; keep
            if level == 1:
                level = 2
            add_heading_custom(out, text, level=level)
        elif text.strip().startswith("→"):
            add_note(out, text)
        elif text.strip().startswith("(1)") or text.strip().startswith("(2)"):
            pe = out.add_paragraph()
            pe.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = pe.add_run(text)
            set_run_font(run, size=11, italic=True)
        else:
            if text.strip():
                add_para(out, text, first_line_indent=True)

    # Append from Referencias onward
    if end_idx is not None:
        for i, p in enumerate(paras):
            if i >= end_idx:
                copy_paragraph(p)

    return out


def main():
    methods = build_methods_doc()
    methods.save(str(OUT_DOWNLOADS))
    methods.save(str(OUT_REPO))
    print(f"Guardado: {OUT_DOWNLOADS}")
    print(f"Guardado: {OUT_REPO}")

    merged = merge_into_intro(methods)
    if merged is not None:
        merged_path = Path("/Users/pablogomez/Downloads/Introduccion_con_Data_y_Metodos_v1.docx")
        merged.save(str(merged_path))
        print(f"Guardado (intro + métodos): {merged_path}")


if __name__ == "__main__":
    main()
