"""Genera Introducción v2 del paper de dry spells, con síntesis bibliográfica."""
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, Cm, Inches
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parent
OUT_REPO = ROOT / "Introduccion_v2_dryspells_chile.docx"
OUT_DOWNLOADS = Path("/Users/pablogomez/Downloads/Introduccion_v2_dryspells_chile.docx")


def set_run_font(run, size=12, bold=False, italic=False):
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic


def add_heading_custom(doc, text, size=14):
    p = doc.add_paragraph()
    run = p.add_run(text)
    set_run_font(run, size=size, bold=True)
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(8)
    return p


def add_body(doc, text):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_after = Pt(8)
    pf.space_before = Pt(0)
    pf.line_spacing = 1.5
    pf.first_line_indent = Inches(0.3)
    pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    run = p.add_run(text)
    set_run_font(run, size=12)
    return p


def add_ref(doc, text):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_after = Pt(4)
    pf.space_before = Pt(0)
    pf.line_spacing = 1.15
    pf.left_indent = Cm(1.25)
    pf.first_line_indent = Cm(-1.25)
    pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    run = p.add_run(text)
    set_run_font(run, size=11)
    return p


def build():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(2.5)
    section.bottom_margin = Cm(2.5)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2.5)

    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(12)

    add_heading_custom(doc, "Introducción")

    add_body(
        doc,
        "Chile central ha atravesado en la última década una megasequía sin precedentes "
        "en el registro instrumental, acompañada de una marcada variabilidad hidrológica "
        "interanual (Garreaud et al., 2017, 2020; Boisier et al., 2016, 2018). Esta "
        "variabilidad plantea un problema que las métricas agregadas de precipitación no "
        "logran capturar por completo: el total anual de lluvia puede ocultar cambios "
        "profundos en la forma en que los períodos secos se distribuyen en el tiempo. Dos "
        "climas con la misma precipitación anual —por ejemplo, 300 mm repartidos en dos "
        "eventos frente a quince— pueden implicar consecuencias muy distintas para "
        "embalses, agricultura y ecosistemas, según cuán prolongadas sean sus rachas secas.",
    )

    add_body(
        doc,
        "Este trabajo aborda esa dimensión de la persistencia de la precipitación, "
        "relevante para la sequía, mediante el estudio de los dry spells: rachas continuas "
        "de días secos, entendidos como días en que la precipitación es menor que un umbral "
        "de día húmedo. En lugar de preguntar cuánto llueve en promedio, el enfoque "
        "distributivo pregunta con qué frecuencia ocurren las rachas y qué tan largas "
        "llegan a ser, incluyendo la cola de eventos extremadamente prolongados. Las "
        "distribuciones de probabilidad (PDFs) de duración de dry spells ofrecen así un "
        "marco parsimonioso para caracterizar la intermitencia de la precipitación, en "
        "línea con el enfoque distributivo aplicado por Martinez-Villalobos y Neelin "
        "(2018, 2019) al estudio de extremos de precipitación, orientado a caracterizar "
        "la cola de las distribuciones mediante escalas de corte en vez de estadísticos "
        "agregados. A esa perspectiva se suma el prototipo de primer paso "
        "(first-passage time) de Stechmann y Neelin (2014): la duración de un spell "
        "húmedo o seco emerge del tiempo que tarda el vapor de agua en cruzar un umbral "
        "de precipitación, lo que da una razón física para estudiar la PDF de duraciones "
        "y no solo su media.",
    )

    add_body(
        doc,
        "La persistencia de períodos secos se ha cuantificado, en la literatura climática, "
        "sobre todo con índices escalares. El más difundido es el de días secos "
        "consecutivos (CDD, del inglés consecutive dry days), definido por el Expert Team "
        "on Climate Change Detection and Indices (ETCCDI) como el máximo anual —o "
        "estacional— de días seguidos con precipitación inferior a 1 mm (Zhang et al., "
        "2011; Sillmann et al., 2013a, 2013b). Otros trabajos reportan el número de días "
        "secos o la duración media de las rachas (Giorgi et al., 2011, 2019; Polade et al., "
        "2014). Un conjunto más reducido examina la distribución completa de duraciones: "
        "Zolina et al. (2013) mostraron que los cambios observados en rachas húmedas y "
        "secas en Europa no se reducen a un cambio en el número de días húmedos, sino a "
        "una reorganización de su agrupación temporal; Lana et al. (2006) ajustaron "
        "distribuciones de probabilidad a dry spells largos en la Península Ibérica. "
        "En Chile, ajustes paramétricos de duración confirman colas pesadas y "
        "desviaciones de un proceso sin memoria (Herrera-Marín et al., 2026). "
        "Estas distinciones importan porque un mismo aumento de días secos puede "
        "fragmentarse en muchas rachas cortas o concentrarse en pocas rachas muy largas. "
        "El CDD retiene el extremo de la cola en cada año, pero no describe con qué "
        "frecuencia ocurren spells de 10, 20 o 50 días ni cómo se desplaza esa cola. Esa "
        "limitación es análoga a la de la precipitación media anual: resume un clima que, "
        "en la práctica, es marcadamente intermitente (Trenberth et al., 2003; "
        "Pendergrass y Knutti, 2018).",
    )

    add_body(
        doc,
        "A este panorama se suman las proyecciones de menor precipitación en Chile central "
        "bajo escenarios de calentamiento, con señal robusta en los ensambles CMIP6 hacia "
        "fines de siglo, incluidas trayectorias de altas emisiones como SSP5-8.5 "
        "(Salazar et al., 2024; Araya-Osses et al., 2020). A escala global, esos mismos "
        "ensambles —y los de CMIP5— proyectan más días secos y dry spells más largos en "
        "subtrópicos y climas mediterráneos, Chile incluido: el aumento de días secos, "
        "más que un cambio de intensidad en los días húmedos, explica gran parte del "
        "secado proyectado (Polade et al., 2014, 2017; Sillmann et al., 2013b; Giorgi "
        "et al., 2019; Wainwright et al., 2021). Saber que “llueve menos en promedio” no "
        "equivale, sin embargo, a saber si las rachas secas se alargan o si aumenta la "
        "probabilidad de spells muy largos. Además, la representación de esa persistencia "
        "en modelos está condicionada por un sesgo sistemático: los modelos climáticos "
        "tienden a llover demasiado a menudo y con demasiado poca intensidad (Stephens "
        "et al., 2010), lo que fragmenta las rachas secas y lleva a subestimar la duración "
        "máxima anual de dry spells (Petrova et al., 2024). Por ello, comparar persistencia "
        "entre observaciones y modelos —o entre productos con distinta frecuencia de días "
        "húmedos— exige controlar esa frecuencia antes de interpretar diferencias de "
        "duración (Martinez-Villalobos et al., 2022). Evaluaciones regionales en "
        "Sudamérica, incluido CORDEX, suelen reportar el CDD o su contraparte húmeda "
        "como un índice anual, no la distribución de duraciones (Skansi et al., 2013; "
        "Reboita et al., 2022).",
    )

    add_body(
        doc,
        "En Chile, el diagnóstico del secado reciente y de la megasequía se ha construido "
        "principalmente a partir de anomalías y tendencias de precipitación y caudales "
        "(Boisier et al., 2016, 2018; Garreaud et al., 2017, 2020), y las proyecciones de "
        "cambio climático se han formulado en términos de acumulados de precipitación y "
        "temperatura, o de impactos hidrológicos asociados (Araya-Osses et al., 2020; "
        "Bozkurt et al., 2018; Salazar et al., 2024). La sequía se caracteriza, en la "
        "práctica operativa y en la literatura hidrológica local, con índices "
        "estandarizados basados en acumulaciones —SPI, SPEI o SSI— cuya escala temporal "
        "se elige según el régimen de la cuenca (Báez-Villanueva et al., 2024). Esos "
        "estudios establecen con claridad que llueve menos y que el déficit tiende a "
        "intensificarse, pero no describen cómo se reorganizan las rachas secas de un día "
        "para otro. El antecedente más cercano es Latoja et al. (2025), quienes, con "
        "CR2MET, documentan en Chile central (32°–40°S) un aumento de la frecuencia de "
        "dry spells junto con una disminución de su duración media —esto es, mayor "
        "intermitencia—, sin evaluar modelos regionales, sin proyectar cambios futuros y "
        "sin caracterizar la cola de duraciones extremas. Herrera-Marín et al. (2026) "
        "modelan, también con CR2MET y umbral de 1 mm, los inicios de dry spells en "
        "invierno austral como un proceso autoexcitado y muestran persistencia con "
        "colas pesadas; Herrera-Marín (2026) descompone el riesgo de inicio bajo "
        "forzante no estacionario. Esos trabajos caracterizan cuándo se inician las "
        "rachas y cómo se agrupan, no cómo cambia la PDF completa de duraciones entre "
        "un clima histórico y uno futuro en un modelo regional. Los caudales, además, "
        "responden a cambios de cobertura de suelo (Escanilla-Minchel et al., 2026): "
        "la persistencia meteorológica de las rachas es un control necesario, pero no "
        "el único, del impacto hidrológico. Lo que permanece abierto es, "
        "por tanto, si un modelo climático regional reproduce la distribución histórica "
        "de duraciones —no solo la media o el máximo anual— una vez igualada la frecuencia "
        "de días húmedos, y si bajo calentamiento aumenta la probabilidad de spells muy "
        "largos, más allá de lo que indica el descenso de la precipitación media.",
    )

    add_body(
        doc,
        "En este estudio preguntamos, primero, en qué medida ALADIN-CLIMATE reproduce la "
        "distribución histórica de duraciones de dry spells respecto de CR2MET, una vez "
        "controladas las diferencias en la frecuencia de días húmedos. Segundo, examinamos "
        "cómo cambia la probabilidad de rachas prolongadas hacia mediados de siglo bajo "
        "SSP5-8.5. De este modo, evaluamos no solo cuánto cambia la precipitación, sino "
        "también cómo se reorganiza temporalmente la persistencia de los períodos secos.",
    )

    add_heading_custom(doc, "Referencias")

    refs = [
        "Araya-Osses, D., Casanueva, A., Román-Figueroa, C., Uribe, J. M., y Paneque, M. (2020). Climate change projections of temperature and precipitation in Chile based on statistical downscaling. Climate Dynamics, 54, 4309–4330. https://doi.org/10.1007/s00382-020-05231-4",
        "Báez-Villanueva, O. M., Zambrano-Bigiarini, M., Miralles, D. G., Beck, H. E., Siegmund, J. F., Alvarez-Garreton, C., Verbist, K., Garreaud, R., Boisier, J. P., y Galleguillos, M. (2024). On the timescale of drought indices for monitoring streamflow drought considering catchment hydrological regimes. Hydrology and Earth System Sciences, 28, 1415–1439. https://doi.org/10.5194/hess-28-1415-2024",
        "Boisier, J. P., Rondanelli, R., Garreaud, R. D., y Muñoz, F. (2016). Anthropogenic and natural contributions to the Southeast Pacific precipitation decline and recent megadrought in central Chile. Geophysical Research Letters, 43(1), 413–421. https://doi.org/10.1002/2015GL067265",
        "Boisier, J. P., Alvarez-Garretón, C., Cordero, R. R., Damiani, A., Gallardo, L., Garreaud, R. D., Lambert, F., Ramallo, C., Rojas, M., y Rondanelli, R. (2018). Anthropogenic drying in central-southern Chile evidenced by long-term observations and climate model simulations. Elementa: Science of the Anthropocene, 6, 74. https://doi.org/10.1525/elementa.328",
        "Bozkurt, D., Rojas, M., Boisier, J. P., y Valdivieso, J. (2018). Projected hydroclimate changes over Andean basins in central Chile from downscaled CMIP5 models under the low and high emission scenarios. Climatic Change, 150, 131–154. https://doi.org/10.1007/s10584-018-2246-7",
        "Escanilla-Minchel, R., Holden, J., y Smith, M. (2026). Land cover change drives river flow modifications in central Chile. Hydrological Processes. https://doi.org/10.1002/hyp.70542",
        "Garreaud, R. D., Alvarez-Garreton, C., Barichivich, J., Boisier, J. P., Christie, D., Galleguillos, M., … Zambrano-Bigiarini, M. (2017). The 2010–2015 megadrought in central Chile: Impacts on regional hydroclimate and vegetation. Hydrology and Earth System Sciences, 21(12), 6307–6327. https://doi.org/10.5194/hess-21-6307-2017",
        "Garreaud, R. D., Boisier, J. P., Rondanelli, R., Damiani, A., Veloso-Aguila, D., y Muñoz, F. (2020). The Central Chile Mega Drought (2010–2018): A climate dynamics perspective. International Journal of Climatology, 40(1), 421–439. https://doi.org/10.1002/joc.6219",
        "Giorgi, F., Im, E.-S., Coppola, E., Diffenbaugh, N. S., Gao, X. J., Mariotti, L., y Shi, Y. (2011). Higher hydroclimatic intensity with global warming. Journal of Climate, 24(20), 5309–5324. https://doi.org/10.1175/2011JCLI3979.1",
        "Giorgi, F., Raffaele, F., y Coppola, E. (2019). The response of precipitation characteristics to global warming from climate projections. Earth System Dynamics, 10, 73–89. https://doi.org/10.5194/esd-10-73-2019",
        "Herrera-Marín, M. (2026). A stochastic hazard decomposition for dry-spell onset dynamics under non-stationary hydroclimatic forcing. Stochastic Environmental Research and Risk Assessment, 40, 172. https://doi.org/10.1007/s00477-026-03309-9",
        "Herrera-Marín, M., Rivera, D., y Godoy-Faúndez, A. (2026). Drought begets drought: self-exciting onsets, spell persistence, and network structure in hydroclimatic extremes. Environmental Research Communications, 8, 051001. https://doi.org/10.1088/2515-7620/ae639a",
        "Lana, X., Martínez, M. D., Burgueño, A., Serra, C., Martín-Vide, J., y Gómez, L. (2006). Distributions of long dry spells in the Iberian Peninsula, years 1951–1990. International Journal of Climatology, 26(14), 1999–2021. https://doi.org/10.1002/joc.1354",
        "Latoja, D., Rivera, D., Lillo-Saavedra, M., y Godoy-Faúndez, Á. (2025). Shifting dry spells in Central Chile: Implications for agricultural water planning. Agricultural & Environmental Letters, 10, e70037. https://doi.org/10.1002/ael2.70037",
        "Martinez-Villalobos, C., y Neelin, J. D. (2018). Shifts in precipitation accumulation extremes during the warm season over the United States. Geophysical Research Letters, 45(16), 8586–8595. https://doi.org/10.1029/2018GL078465",
        "Martinez-Villalobos, C., y Neelin, J. D. (2019). Why do precipitation intensities tend to follow gamma distributions? Journal of the Atmospheric Sciences, 76(11), 3611–3631. https://doi.org/10.1175/JAS-D-18-0343.1",
        "Martinez-Villalobos, C., Neelin, J. D., y Pendergrass, A. G. (2022). Metrics for evaluating CMIP6 representation of daily precipitation probability distributions. Journal of Climate, 35(17), 5719–5743. https://doi.org/10.1175/JCLI-D-21-0617.1",
        "Pendergrass, A. G., y Knutti, R. (2018). The uneven nature of daily precipitation and its change. Geophysical Research Letters, 45(21), 11980–11988. https://doi.org/10.1029/2018GL080298",
        "Petrova, I. Y., Miralles, D. G., Brient, F., Donat, M. G., Min, S.-K., Kim, Y.-H., y Bador, M. (2024). Observation-constrained projections reveal longer-than-expected dry spells. Nature, 633, 594–600. https://doi.org/10.1038/s41586-024-07887-y",
        "Polade, S. D., Pierce, D. W., Cayan, D. R., Gershunov, A., y Dettinger, M. D. (2014). The key role of dry days in changing regional climate and precipitation regimes. Scientific Reports, 4, 4364. https://doi.org/10.1038/srep04364",
        "Polade, S. D., Gershunov, A., Cayan, D. R., Dettinger, M. D., y Pierce, D. W. (2017). Precipitation in a warming world: Assessing projected hydro-climate changes in California and other Mediterranean climate regions. Scientific Reports, 7, 10783. https://doi.org/10.1038/s41598-017-11285-y",
        "Reboita, M. S., da Rocha, R. P., de Souza, C. A., Baldoni, T. C., Silva, P. L. L. da S., y Ferreira, G. W. S. (2022). Future projections of extreme precipitation climate indices over South America based on CORDEX-CORE multimodel ensemble. Atmosphere, 13(9), 1463. https://doi.org/10.3390/atmos13091463",
        "Salazar, Á., Thatcher, M., Goubanova, K., Bernal, P., Gutiérrez, J., y Squeo, F. (2024). CMIP6 precipitation and temperature projections for Chile. Climate Dynamics. https://doi.org/10.1007/s00382-023-07034-9",
        "Sillmann, J., Kharin, V. V., Zhang, X., Zwiers, F. W., y Bronaugh, D. (2013a). Climate extremes indices in the CMIP5 multimodel ensemble: Part 1. Model evaluation in the present climate. Journal of Geophysical Research: Atmospheres, 118, 1716–1733. https://doi.org/10.1002/jgrd.50203",
        "Sillmann, J., Kharin, V. V., Zwiers, F. W., Zhang, X., y Bronaugh, D. (2013b). Climate extremes indices in the CMIP5 multimodel ensemble: Part 2. Future climate projections. Journal of Geophysical Research: Atmospheres, 118, 2473–2493. https://doi.org/10.1002/jgrd.50188",
        "Skansi, M. de los M., Brunet, M., Sigró, J., Aguilar, E., Groening, J. A. A., Bentancur, O. J., … Jones, P. D. (2013). Warming and wetting signals emerging from analysis of changes in climate extreme indices over South America. Global and Planetary Change, 100, 295–307. https://doi.org/10.1016/j.gloplacha.2012.11.004",
        "Stechmann, S. N., y Neelin, J. D. (2014). First-passage-time prototypes for precipitation statistics. Journal of the Atmospheric Sciences, 71(9), 3269–3291. https://doi.org/10.1175/JAS-D-13-0268.1",
        "Stephens, G. L., L’Ecuyer, T., Forbes, R., Gettelman, A., Golaz, J.-C., Bodas-Salcedo, A., Suzuki, K., Gabriel, P., y Haynes, J. (2010). Dreary state of precipitation in global models. Journal of Geophysical Research: Atmospheres, 115, D24211. https://doi.org/10.1029/2010JD014532",
        "Trenberth, K. E., Dai, A., Rasmussen, R. M., y Parsons, D. B. (2003). The changing character of precipitation. Bulletin of the American Meteorological Society, 84(9), 1205–1218. https://doi.org/10.1175/BAMS-84-9-1205",
        "Wainwright, C. M., Black, E., y Allan, R. P. (2021). Future changes in wet and dry season characteristics in CMIP5 and CMIP6 simulations. Journal of Hydrometeorology, 22, 2339–2357. https://doi.org/10.1175/JHM-D-21-0017.1",
        "Zhang, X., Alexander, L., Hegerl, G. C., Jones, P., Klein Tank, A., Peterson, T. C., Trewin, B., y Zwiers, F. W. (2011). Indices for monitoring changes in extremes based on daily temperature and precipitation data. Wiley Interdisciplinary Reviews: Climate Change, 2(6), 851–870. https://doi.org/10.1002/wcc.147",
        "Zolina, O., Simmer, C., Belyaev, K., Gulev, S. K., y Koltermann, P. (2013). Changes in the duration of European wet and dry spells during the last 60 years. Journal of Climate, 26(6), 2022–2047. https://doi.org/10.1175/JCLI-D-11-00498.1",
    ]
    for r in refs:
        add_ref(doc, r)

    doc.save(OUT_REPO)
    try:
        doc.save(OUT_DOWNLOADS)
    except OSError:
        pass
    return OUT_REPO


if __name__ == "__main__":
    path = build()
    print("Wrote", path)
