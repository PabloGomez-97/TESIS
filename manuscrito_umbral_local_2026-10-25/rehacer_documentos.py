from pathlib import Path
from copy import deepcopy
import csv,json,hashlib,shutil,re
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH

ROOT=Path(__file__).resolve().parent.parent
OUT=Path(__file__).resolve().parent
A=ROOT/'temporal_umbral_por_celda_2026-10-24/ejecucion'
B=ROOT/'temporal_umbral_por_celda_0p1mm_2026-10-25/ejecucion'
COMP=ROOT/'temporal_umbral_por_celda_0p1mm_2026-10-25/comparativa'
def rows(p):
 with open(p) as f:return list(csv.DictReader(f))
def fmt(x,n=1):return f'{float(x):.{n}f}'.replace('.',',')
def integer(x):return f'{int(float(x)):,}'.replace(',',' ')
def setp(p,text):
 rpr=deepcopy(p.runs[0]._r.rPr) if p.runs and p.runs[0]._r.rPr is not None else None
 p.clear();r=p.add_run(text)
 if rpr is not None:r._r.insert(0,rpr)
 for tag in ['highlight','shd','color']:
  for e in r._r.findall('.//'+qn('w:'+tag)):e.getparent().remove(e)
 return p
def clean(d):
 for s in d.styles:
  if s.type==1 and (s.name.startswith('Heading') or s.name=='Title'):s.font.color.rgb=RGBColor(0,0,0)
 for p in list(d.paragraphs)+[p for t in d.tables for r in t.rows for c in r.cells for p in c.paragraphs]:
  for run in p.runs:
   if run.font.name=='Gungsuh':run.font.name='Times New Roman'
  for e in p._p.findall('.//'+qn('w:highlight')):e.getparent().remove(e)
  if p.style.name.startswith('Heading'):
   for r in p.runs:r.font.color.rgb=RGBColor(0,0,0)
 for p in list(d.paragraphs):
  if not p.text.strip() and not p._p.findall('.//'+qn('w:drawing')):p._p.getparent().remove(p._p)
  if re.match(r'^\d+(?:\.\d+)*\.? ',p.text) or re.match(r'^S\d+\.',p.text):p.paragraph_format.keep_with_next=True
  if p.text.startswith(('Figura ','Figura S')):p.paragraph_format.keep_together=True
  if p._p.findall('.//'+qn('w:drawing')):p.paragraph_format.keep_with_next=True
 for t in d.tables:
  t.autofit=False
  grid=t._tbl.tblGrid.gridCol_lst
  old=[int(round(float(e.get(qn('w:w'))))) for e in grid]
  widths=[int(9360*v/sum(old)) for v in old]
  for e,w in zip(grid,widths):e.set(qn('w:w'),str(w))
  tw=t._tbl.tblPr.find(qn('w:tblW'))
  if tw is not None:tw.set(qn('w:w'),'9360');tw.set(qn('w:type'),'dxa')
  for r in t.rows:
   trpr=r._tr.get_or_add_trPr();trpr.append(OxmlElement('w:cantSplit'))
   for c,w in zip(r.cells,widths):
    c.width=int(w*635)
    for p in c.paragraphs:
     p.paragraph_format.space_after=Pt(3);p.paragraph_format.space_before=Pt(3)
     p.paragraph_format.line_spacing=1;p.alignment=WD_ALIGN_PARAGRAPH.LEFT
     for run in p.runs:run.font.size=Pt(9)
  trpr=t.rows[0]._tr.get_or_add_trPr();repeat=OxmlElement('w:tblHeader');trpr.append(repeat)
 for sh in d.inline_shapes:
  ratio=sh.height/sh.width
  width=min(sh.width, Inches(6.5), int(Inches(6.2)/ratio));sh.width=width;sh.height=int(width*ratio)
def table_fill(t,data):
 for e in t._tbl.findall('.//'+qn('w:gridCol')):
  if e.get(qn('w:w')):e.set(qn('w:w'),str(round(float(e.get(qn('w:w'))))))
 while len(t.rows)>1:t._tbl.remove(t.rows[-1]._tr)
 for vals in data:
  cells=t.add_row().cells
  for c,v in zip(cells,vals):setp(c.paragraphs[0],str(v))
def new_table(d,headers,data):
 t=d.add_table(rows=1,cols=len(headers))
 t._tbl.remove(t._tbl.tblPr);t._tbl.insert(0,deepcopy(d.tables[0]._tbl.tblPr))
 for c,h in zip(t.rows[0].cells,headers):c.text=h
 table_fill(t,data);return t
FIG=[]
def copyfig(src,name):
 dst=OUT/'figuras'/name
 if name not in ['Figura_05.png','Figura_06.png','Figura_S3.png','Figura_S4.png']:shutil.copy2(src,dst)
 assert dst.exists(),f'Ejecutar generar_rr_publicacion.py antes de crear {name}'
 FIG.append({'archivo':name,'fuente':str(src.relative_to(ROOT)),'redibujado':name in ['Figura_05.png','Figura_06.png','Figura_S3.png','Figura_S4.png'],'sha256':hashlib.sha256(dst.read_bytes()).hexdigest()});return dst
def replace_images(d,paths):
 for sh,path in zip(d.inline_shapes,paths):
  embed=sh._inline.xpath('.//a:blip')[0].get(qn('r:embed'))
  part=d.part.related_parts[embed];part._blob=path.read_bytes()
  from PIL import Image
  with Image.open(path) as im:w,h=im.size
  width=min(Inches(6.5),int(Inches(7.0)*w/h));sh.width=width;sh.height=int(width*h/w)
def add_fig(d,path,caption):
 from PIL import Image
 with Image.open(path) as im:w,h=im.size
 p=d.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
 p.add_run().add_picture(str(path),width=min(Inches(6.5),int(Inches(7)*w/h)))
 d.add_paragraph(caption)

paper_src=next((OUT/'fuentes').glob('*v1 (3).docx'))
supp_src=OUT/'fuentes/Material suplementario.docx'
d=Document(paper_src);p=list(d.paragraphs)
updates={
0:'Material suplementario de la versión con calibración local.',
1:'',2:'',
4:'Duración de las rachas secas a lo largo de un gradiente climático de Chile: evaluación de un modelo regional y cambio proyectado a mediados de siglo bajo SSP5-8.5',
14:'Este trabajo estudia la persistencia mediante dry spells: rachas consecutivas de días con precipitación inferior a un umbral de día húmedo. El enfoque distributivo caracteriza cuántas rachas ocurren, cuánto duran y cómo se comporta su cola. Para representar simultáneamente su número y sus duraciones se utilizan histogramas divididos por el ancho de bin, cuya integral corresponde al número total de rachas; no son densidades normalizadas a uno.',
16:'A ese panorama se suman las proyecciones de menor precipitación en Chile central bajo calentamiento (Araya-Osses et al., 2020; Salazar et al., 2024). A escala global, el aumento del número de días secos contribuye al secado proyectado en numerosos climas subtropicales y mediterráneos (Polade et al., 2014). Saber que llueve menos en promedio no determina, sin embargo, cómo cambia la distribución de duraciones de las rachas. La pregunta central es cómo se modifica esa distribución y, en especial, el número de rachas prolongadas.',
18:'Responder esa pregunta con un modelo regional exige evaluar primero su representación histórica. Una frecuencia excesiva de días húmedos puede interrumpir las rachas y acortarlas. La evaluación de distribuciones de duración con umbrales dependientes del modelo permite reducir diferencias en la frecuencia de días húmedos antes de comparar las rachas (Leung et al., 2022, §3.b). Petrova et al. (2024) muestran la importancia de los sesgos de la racha seca anual más larga para interpretar proyecciones; su métrica anual y su exclusión de climas con rachas históricas superiores a 300 días difieren del análisis de todas las rachas completas que realizamos aquí.',
34:'Se define wet day como un día con precipitación pr ≥ R y dry day como un día con pr < R. Un dry spell es una secuencia consecutiva de dry days; su duración t es el número de días de la secuencia (Figura 2). En CR2MET se adopta R = 1 mm día⁻¹, siguiendo la convención de día húmedo del ETCCDI (Zhang et al., 2011) y el antecedente de Latoja et al. (2025) en Chile central. Esta elección favorece la comparabilidad, pero no constituye un umbral universal de eficacia hidrológica. La referencia de 0,1 mm también es válida en la literatura y se evalúa como sensibilidad a la precipitación débil (S8). En cada período se identifican todas las rachas de la serie diaria completa y se excluyen las que tocan su primer o último día, cuya duración total se desconoce. Las rachas censuradas representan menos del 1,1 % en Loa y del 0,23 % en las demás ventanas, considerando los tres conjuntos de la Tabla 3.',
39:'Tabla 2. Criterios de umbral. Se conservan las etiquetas i–iii de los experimentos de sensibilidad; el análisis principal utiliza iii, calibración por celda.',
41:'Con 1 mm día⁻¹ en ambos productos, la fracción espacial media de días húmedos es 30,685 % en ALADIN y 20,945 % en CR2MET. El análisis principal utiliza un umbral histórico R₀(x,y) por celda de ALADIN, calibrado contra CR2MET previamente interpolado. Para cada celda se resuelve F_ALADIN,x,y(R₀) = F_CR2MET,x,y(1 mm), donde F es la fracción de días húmedos de todos los días de 1980–2014. Se usa bisección con tolerancia de 10⁻⁴ en fracción, equivalente a 0,01 puntos porcentuales. El intervalo inicial es 0–15 mm día⁻¹; en las 82 celdas que requieren valores mayores se amplía hasta superar el máximo histórico, evitando dejar huecos por una cota numérica arbitraria. Las 3 587 celdas quedan calibradas, con error absoluto inferior a 0,008 puntos porcentuales. Los umbrales tienen mediana de 6,489 mm día⁻¹, p10 de 2,285 y p90 de 11,602; el rango es 0,039–25,026 mm día⁻¹ (Figura 3 y Figura S2). El campo histórico se conserva sin recalibrar en 2040–2074 y en ambos escenarios. La calibración ajusta una frecuencia marginal anual por celda; no ajusta duraciones ni secuencias temporales. El umbral global de 5,285 mm día⁻¹ y el umbral fijo de 1 mm en ALADIN quedan como sensibilidades (criterios i y ii). El uso de umbrales dependientes del modelo tiene antecedentes en Leung et al. (2022), aunque esa referencia no valida por sí sola esta implementación espacial.',
43:'Figura 3. Calibración histórica por celda con CR2MET a 1 mm día⁻¹. Distribución del campo R₀(x,y) y comparación de las frecuencias de días húmedos entre productos antes y después de calibrar. La igualdad se impone por celda sobre 1980–2014, dentro de una tolerancia de 0,01 puntos porcentuales. El campo se mantiene fijo para las proyecciones.',
44:'R₀(x,y) es un umbral estadístico de clasificación: sus valores no son equivalentes físicos de 1 mm observado ni corrigen la cantidad precipitada. Igualar la frecuencia anual por celda evita que el clima de otras zonas determine su umbral, pero no impone la frecuencia de cada mes ni la de marzo–noviembre. Tampoco determina el orden de los días húmedos. Por ello, las diferencias de duración que persisten son discrepancias de la distribución temporal después de controlar esa frecuencia anual; no identifican por sí solas un mecanismo de circulación ni una causa única del sesgo.',
55:'donde N(t ≥ D) es el número de rachas completas de al menos D días en ALADIN futuro e histórico, con el mismo campo histórico R₀(x,y). Cada racha se contabiliza por celda y los conteos se agregan sobre la región. Al comparar períodos de igual duración, RR > 1 indica más excedencias futuras. No es el cociente de las fracciones de rachas largas ni una probabilidad individual de sequía.',
56:'Los risk ratios se calculan en las cuatro regiones administrativas considerando rachas cuyo inicio cae entre marzo y noviembre, con su duración completa aunque se extiendan al verano. En CR2MET esos meses concentran entre el 85,1 % y el 97,3 % de la precipitación anual de estas regiones, y entre el 81,6 % y el 95,9 % en ALADIN. Las distribuciones por ventana incluyen todo el año. Se evalúan D = 5, 10, 15, 20, 25, 30, 35, 40, 50 y 60 días; D = 20 es una referencia descriptiva y no un umbral validado de daño agrícola. Los IC95 se estiman con 1 000 réplicas de bootstrap de 35 años de inicio, sorteados de forma independiente para cada período y conjuntamente para todas las celdas de la región, conservando su dependencia espacial dentro de cada año. Las réplicas con N_hist* = 0 se cuentan y descartan; el intervalo queda condicionado a las restantes. A D = 20 no hay descartes en el análisis principal. Se considera significativo un IC95 completamente sobre o bajo 1. Los intervalos son puntuales, sin corrección por comparaciones múltiples entre duraciones o regiones; se interpreta la curva completa y se evita atribuir robustez a un cruce aislado. El bootstrap no recalibra R₀: sus intervalos son condicionales al campo estimado.',
58:'Se comparan los criterios i global, ii fijo y iii local, conservando CR2MET a 1 mm, y se examina la sensibilidad al mínimo de rachas por celda, al regrillado y al escenario SSP3-7.0 (S1–S7). Además, se repite la cadena con CR2MET a 0,1 mm y un campo R₀(x,y) calibrado específicamente para esa referencia, también fijo en el futuro (S8). Esta última prueba modifica tanto la definición observacional como el umbral del modelo; no equivale a imponer 0,1 mm en ALADIN. Se incluyen diagnósticos de frecuencia seca, persistencia de primer orden y una referencia mensual de Markov (S7–S9).',
63:'La primera pregunta es si ALADIN reproduce las duraciones históricas después de igualar la frecuencia anual de días húmedos por celda, con CR2MET a 1 mm y R₀(x,y) en ALADIN (criterio iii).',
68:'Con la calibración por celda, ALADIN reproduce el gradiente espacial, pero las duraciones siguen siendo menores en promedio. En las 3 549 celdas comunes con al menos 30 rachas en ambos productos, la media espacial de las duraciones medias es 42,1 días en CR2MET y 35,3 en ALADIN: Δ = −6,7 días. Para t99, los promedios son 221,1 y 185,4 días: Δ = −35,7 días. Los sesgos medianos son −1,9 y −6,3 días, respectivamente. ALADIN es significativamente más corto en el 49,1 % de las celdas para la media y el 25,4 % para t99, frente a un 0,2 % y un 4,9 % donde es significativamente más largo (Figura 4; Tabla S1). Son promedios de estadísticos por celda, distintos de los estadísticos de las rachas agrupadas por ventana.',
69:'Con 1 mm fijo en ambos productos, los sesgos medios son −27,2 días en duración media y −134,9 días en t99. El umbral global reduce esos valores, pero no iguala la frecuencia dentro de cada celda. En una comparación de ambos métodos sobre las mismas 3 526 celdas, el error absoluto medio de la duración media disminuye de 14,890 a 6,673 días y el de t99 de 71,473 a 48,825 días (S1). La mejora se evalúa sobre una máscara común y no elimina el sesgo residual ni garantiza la representación de la cola.',
71:'Figura 4. Duración histórica de las rachas secas con calibración local, 1980–2014. Filas: duración media y t99. Columnas: CR2MET a 1 mm, ALADIN con R₀(x,y) y Δ = ALADIN − CR2MET. Los seis paneles muestran las mismas 3 549 celdas con al menos 30 rachas en ambos productos.',
74:'Las distribuciones agrupadas muestran un sesgo dependiente de la ventana y de la métrica (Figura 5; Tabla 3). En Loa, ALADIN tiene media de 50,5 días frente a 65,6 en CR2MET (−23,0 %) y p99 de 353 frente a 456 días (−22,6 %). En Maipo la media es 10,4 frente a 11,0 días (−4,8 %), pero el p99 es 121 frente a 119 días (+1,7 %). En Maule la media es 8,2 frente a 9,5 días (−13,6 %) y p99 de 67 frente a 83 (−19,3 %); en Biobío la media es 5,4 frente a 6,2 (−12,6 %) y p99 de 35 frente a 44 (−20,5 %). Por tanto, no hay una subestimación del p99 en las cuatro ventanas: Maipo constituye una excepción tras la calibración local.',
75:'El número de rachas sigue siendo mayor en ALADIN: por ejemplo, 71 662 frente a 55 015 en Loa y 157 567 frente a 150 209 en Maipo. Estas diferencias persisten con la frecuencia anual prácticamente igualada por celda. La calibración controla el número de días húmedos, pero no el número de interrupciones independientes ni las duraciones de las secuencias secas y húmedas.',
76:'Tabla 3. Estadísticos de las rachas completas agrupadas por ventana con calibración local. Histórico 1980–2014 y futuro SSP5-8.5, 2040–2074. Duraciones en días; la última columna es el porcentaje de rachas de al menos 20 días. Son resultados descriptivos, sin intervalos de confianza del cambio futuro.',
78:'Figura 5. Distribuciones empíricas de duraciones con calibración local. Negro: CR2MET histórico; azul: ALADIN histórico; rojo discontinuo: ALADIN futuro SSP5-8.5. Paneles: Loa, Maipo, Maule y Biobío. Ejes logarítmicos y bins comunes dentro de cada ventana; la integral es n_spells, no uno. El p99 histórico de ALADIN es menor en Loa, Maule y Biobío y ligeramente mayor en Maipo.',
80:'La segunda pregunta es cómo cambia la ocurrencia de rachas prolongadas en esta simulación bajo SSP5-8.5. Se comparan ALADIN histórico y futuro utilizando el mismo campo R₀(x,y), sin imponer que la frecuencia húmeda futura coincida con la histórica.',
82:'A D = 20 días, solo La Araucanía presenta un aumento significativo: de 2 185 a 3 789 rachas, RR = 1,7341, IC95 [1,0803; 3,0445]. Coquimbo pasa de 26 613 a 28 788, RR = 1,0817 [0,9859; 1,2040]; O’Higgins, de 8 202 a 9 511, RR = 1,1596 [0,9735; 1,3864]; y Los Lagos, de 328 a 371, RR = 1,1311 [0,3255; 8,1138]. Los tres últimos intervalos incluyen 1. Un cociente mayor que 1 no basta para concluir un aumento estadísticamente detectable (Tabla 4; Figura 6).',
83:'La inferencia depende de la definición y del escenario. Con CR2MET a 0,1 mm y recalibración por celda, Coquimbo presenta RR(20) = 1,1228 [1,0101; 1,2497] y La Araucanía 1,9236 [1,1413; 3,5629], ambos significativos; O’Higgins y Los Lagos siguen con intervalos que incluyen 1. El cruce del límite de significancia en Coquimbo muestra sensibilidad a la precipitación débil, pero no prueba una diferencia significativa entre los dos RR. Con SSP3-7.0 y referencia de 1 mm, La Araucanía queda marginalmente sobre 1, RR = 1,5547 [1,0037; 2,5920], y las otras tres regiones incluyen 1. Con la referencia de 0,1 mm, ningún RR(20) es significativo bajo SSP3-7.0 (S6 y S8).',
84:'Tabla 4. RR(20) de conteos de rachas con inicio marzo–noviembre, ALADIN con calibración local histórica contra CR2MET a 1 mm, SSP5-8.5. IC95 por bootstrap de años, 1 000 réplicas; ninguna descartada a D = 20.',
89:'Figura 6. Curvas RR(D) de ALADIN SSP5-8.5 respecto del histórico, con R₀(x,y) fijo y rachas iniciadas entre marzo y noviembre. Línea horizontal: RR = 1; bandas: IC95 puntuales. Los símbolos huecos señalan duraciones con réplicas descartadas por N_hist* = 0, donde no se representa banda. A 20 días solo La Araucanía tiene el IC95 completamente sobre 1. Si el conteo histórico es cero, el RR no está definido.',
91:'En las cuatro ventanas disminuye el número de rachas y aumenta su duración media (Tabla 3). En Loa, la media pasa de 50,5 a 52,2 días y el p99 de 353 a 368,5; en Maipo, de 10,4 a 12,0 y de 121 a 138; en Maule, de 8,2 a 9,6 y de 67 a 89; y en Biobío, de 5,4 a 6,1 y de 35 a 41. Son desplazamientos descriptivos de estas simulaciones; no se ha estimado la significancia del cambio de cada estadístico por ventana. La censura excluye algunas rachas potencialmente prolongadas, pero la duración completa de esas rachas no se conoce y no permite cuantificar una corrección exacta de la cola.',
92:'El máximo muestral es ruidoso y no resume por sí solo la distribución. En Biobío disminuye de 153 a 145 días, aunque aumentan la media y el p99 (Tabla 3). Por eso se examinan conjuntamente las duraciones, el número de rachas y las curvas de excedencias.',
93:'En conjunto, las distribuciones describen menos rachas y mayores duraciones medias en las cuatro ventanas, con extensión del p99 particularmente marcada en Maipo, Maule y Biobío. En las regiones administrativas, el aumento del conteo de rachas de al menos 20 días solo se detecta en La Araucanía con la definición principal. Los Lagos mantiene una incertidumbre amplia por sus pocos eventos largos.',
96:'La calibración por celda permite evaluar las duraciones después de igualar la frecuencia anual húmeda local. ALADIN reproduce el gradiente norte–sur, pero conserva un sesgo medio negativo en duración y t99. Esto indica que controlar una frecuencia marginal no determina la distribución de rachas. La mejora respecto del umbral global es importante para la comparación histórica, aunque no transforma el campo de precipitación en observaciones corregidas. El contraste con Leung et al. (2022), quienes encuentran rachas más largas en los modelos de su evaluación, muestra que el signo del sesgo no es universal. Petrova et al. (2024) estudian extremos anuales y no validan directamente la cola de todas las rachas de nuestro norte hiperárido.',
97:'Hacia mediados de siglo, las cuatro ventanas muestran menos rachas y duraciones medias mayores. La frecuencia de días húmedos cae un 5,6 % en el dominio y entre un 10,2 % y un 12,9 % en Maipo, Maule y Biobío; en Loa aumenta un 1,1 %. En las cuatro regiones administrativas aumenta la frecuencia seca de marzo–noviembre entre 1,6 y 3,9 puntos porcentuales, con IC95 sobre cero. Los intervalos de Δρ incluyen cero en todas ellas con la referencia de 1 mm, aunque en Los Lagos el límite inferior es −0,00068 (S7). Una referencia mensual de Markov que modifica la frecuencia y conserva la persistencia histórica produce RR compatibles con los IC95 observados en los 30 casos con al menos 100 eventos históricos (S9). Esto es compatibilidad con una explicación basada en menor frecuencia húmeda, no una atribución causal ni prueba de ausencia de cambios en dependencia temporal. Con 0,1 mm, Los Lagos sí presenta un aumento detectable de ρ (S8). El contraste con las rachas históricas más numerosas y cortas de Latoja et al. (2025) debe considerar sus períodos, estaciones y reglas de conteo diferentes. Un único miembro no separa señal forzada y variabilidad interna.',
99:'La evaluación histórica muestra que comparar precipitación media o frecuencia húmeda no basta para evaluar las rachas secas. Conviene comparar también la distribución de duraciones después de controlar la frecuencia local, declarando su definición y su sensibilidad (Leung et al., 2022). La cercanía del p99 en Maipo no se extiende a Loa, Maule y Biobío. El sesgo histórico limita la confianza en las magnitudes futuras y no permite corregirlas aplicando automáticamente una relación obtenida para máximos anuales en otros modelos.',
100:'Para la gestión hídrica, las curvas de excedencias aportan información sobre cuántas rachas superan diferentes duraciones. Este trabajo no estima impactos agrícolas, caudales ni daños: una racha bajo 1 mm no equivale necesariamente a déficit de humedad del suelo, y D = 20 días no es un umbral de impacto validado.',
102:'Varias limitaciones acotan la interpretación. Se utiliza un RCM, un miembro por escenario y un escenario principal; los resultados no representan incertidumbre multimodelo. Se supone transferible al futuro el campo histórico de umbrales, y el bootstrap no incluye incertidumbre de esa calibración. Igualar frecuencia anual no iguala su estacionalidad ni las cantidades de lluvia. El umbral es una clasificación estadística: ALADIN conserva sesgos importantes en los acumulados. Los pools contienen celdas dependientes espacialmente; el remuestreo conjunto por años conserva esa dependencia dentro de cada año, pero no explora explícitamente bloques multianuales. En los mapas se excluyen 38 de 3 587 celdas por el mínimo de 30 rachas; los extremos del norte y las duraciones t99 son sensibles al tamaño muestral (S1). La censura de bordes omite eventos de duración total desconocida (S9). En Los Lagos los pocos eventos largos generan intervalos amplios y réplicas sin denominador a ciertas duraciones. Los intervalos son puntuales y no se corrigen por multiplicidad. Las diferencias entre las referencias de 1 y 0,1 mm, especialmente en Coquimbo y en la persistencia de Los Lagos, limitan la robustez de afirmaciones independientes del umbral (S8). El estudio no diagnostica mecanismos de circulación ni impactos hidrológicos.',
104:'Bajo SSP5-8.5, esta simulación proyecta menos rachas y mayores duraciones medias en las cuatro ventanas estudiadas, con extensión del p99 en todas ellas. En los conteos regionales de rachas de al menos 20 días, el aumento es significativo solo en La Araucanía con la referencia principal de 1 mm. El resultado de Coquimbo depende de la definición de día húmedo. Caracterizar todas las rachas completas añade información a un único extremo anual, y la calibración por celda mejora la evaluación histórica sin eliminar sus limitaciones.'
}
for i,t in updates.items():setp(p[i],t)
setp(p[44],p[44].text+' Se calibra la frecuencia para conservar las duraciones como variable de evaluación independiente del ajuste. Calibrar directamente contra ellas sería otro objetivo de corrección y requeriría validación independiente de las duraciones ajustadas.')
setp(p[51],p[51].text.replace('y una celda con menos de 50 réplicas válidas se considera no significativa','y no se declara significancia en una celda con menos de 50 réplicas válidas; esta es una salvaguarda numérica, no evidencia de igualdad entre productos'))
setp(p[17],p[17].text.replace('Herrera-Marín (2026) y Herrera-Marín et al. (2026) modelan','Herrera-Marín et al. (2026) modelan'))
p[0].runs[0].font.size=Pt(11);p[0].runs[0].font.bold=False
p[4].style=d.styles['Title']
# Conservar estructura y antecedentes, resolver citas decorativas conocidas.
setp(p[15], 'La persistencia seca suele resumirse con índices escalares. El CDD retiene el máximo anual de días consecutivos con precipitación inferior a 1 mm (Zhang et al., 2011); no describe la frecuencia de todas las duraciones. Zolina et al. (2013) examinan distribuciones de rachas europeas y distinguen cambios en el número de días húmedos y en su agrupación. En Chile, Herrera-Marín et al. (2026) describen persistencia y colas de duración con CR2MET. Estas distinciones motivan examinar todas las rachas completas y separar frecuencia marginal y dependencia temporal.')
setp(p[66],'El gradiente aparece tanto en los estadísticos por celda como en las distribuciones agrupadas. Una evaluación regional debe conservar esa heterogeneidad: las rachas de cientos de días del norte requieren reglas que permitan cruzar el año calendario, mientras que en el sur los eventos largos son mucho menos frecuentes.')
table_fill(d.tables[0],[['CR2MET v2.5','Referencia observacional grillada','Lat/lon 0,05°; interpolada a ALADIN','1980–2014','1 mm día⁻¹'],['ALADIN CHP12 histórico','RCM CNRM-ALADIN64; CNRM-ESM2-1','Nativa ~12 km; 3 587 celdas','1980–2014','R₀(x,y), criterio iii'],['ALADIN SSP5-8.5','Proyección principal','Idéntica al histórico','2040–2074','Campo histórico fijo'],['ALADIN SSP3-7.0','Sensibilidad de RR (S6)','Idéntica al histórico','2040–2074','Campo histórico fijo']])
table_fill(d.tables[1],[['i (global)','1 mm','5,285 mm, único','Sensibilidad espacial'],['ii (fijo)','1 mm','1 mm','Sensibilidad sin calibrar'],['iii (local)','1 mm','R₀(x,y) por celda','Análisis principal'],['Referencia 0,1 mm','0,1 mm','Campo propio R₀,₀.₁(x,y)','Sensibilidad S8']])
t3=rows(A/'resultados_tesis/tabla3_estadisticos_ventanas.csv')
table_fill(d.tables[2],[[r['ventana'],r['producto'],integer(r['n_spells']),fmt(r['media']),fmt(r['mediana']),fmt(r['p90']),fmt(r['p99']),integer(r['max']),fmt(r['frac_ge20_pct'])] for r in t3])
rr=json.loads((A/'outputs_pregunta5_aladin/results.json').read_text())['rr_table']
table_fill(d.tables[3],[[r['Region'],integer(r['N_hist']),integer(r['N_fut']),fmt(r['RR observado'],4),f"[{fmt(r['IC95 inferior'],4)}; {fmt(r['IC95 superior'],4)}]",'Sí' if r['Sig aumento'] else 'No'] for r in rr])
mainpaths=[copyfig(A/'paper_figuras'/n,f'Figura_{i:02d}.png') for i,n in enumerate(['fig02_dominio_cuencas_regiones.png','fig_esquema_dryspell.png','fig03_calibracion_R0_ftau.png','fig05_mapas_persistencia_criterio_i.png','fig42_pdf_cuencas_loglog_panel.png','fig08_rr_curves_aladin.png'],1)]
replace_images(d,mainpaths)
# Eliminar referencias no citadas tras resolver las notas del borrador.
body='\n'.join(x.text for x in d.paragraphs[:107])
remove=[116,117,118,120,125,126,127,128,131,133,134,135,136,137]
for i in remove:p[i]._p.getparent().remove(p[i]._p)
# Metadatos completos ya verificados en el registro bibliográfico local.
setp(p[122], 'Leung, L. R., Boos, W. R., Catto, J. L., DeMott, C. A., Martin, G. M., Neelin, J. D., O’Brien, T. A., Xie, S., Feng, Z., Klingaman, N. P., Kuo, Y.-H., Lee, R. W., Martinez-Villalobos, C., Vishnu, S., Priestley, M. D. K., Tao, C., y Zhou, Y. (2022). Exploratory precipitation metrics: Spatiotemporal characteristics, process-oriented, and phenomena-based. Journal of Climate, 35(12), 3659–3686. https://doi.org/10.1175/JCLI-D-21-0590.1')
for i in [123,124]:p[i]._p.getparent().remove(p[i]._p)
for para in d.paragraphs:
 text=para.text.replace('por cuenca','por ventana').replace('En cada cuenca','En cada ventana').replace('según la cuenca','según la ventana')
 if text!=para.text:setp(para,text)
clean(d);d.save(OUT/'Paper_umbral_local_1mm.docx')

s=Document(supp_src);p=list(s.paragraphs)
su={
1:'Este material complementa §2.7 del artículo. Mantiene las sensibilidades históricas al umbral (S1–S2), documenta el campo local (S3), el umbral fijo de ALADIN (S4), el regrillado (S5), SSP3-7.0 (S6) y la frecuencia seca y persistencia (S7). Se añade la referencia de 0,1 mm con recalibración por celda (S8) y se explicitan los diagnósticos Markov y de censura (S9). Salvo indicación contraria, la referencia es CR2MET a 1 mm y el análisis principal es ALADIN con umbral por celda, 1980–2014 y 2040–2074, SSP5-8.5. Se excluyen rachas que tocan los bordes de cada período. Las ventanas son cajas geográficas, no polígonos oficiales de cuenca.',
3:'La Tabla S1 conserva las etiquetas del experimento: i global, ii fijo y iii local, este último principal. Los promedios usan celdas con al menos 30 rachas en ambos productos dentro de cada criterio. Las máscaras pueden diferir entre criterios; no se debe atribuir toda diferencia de sus promedios a la calibración. La significancia se estima con 1 000 remuestreos anuales independientes entre productos, condicionales a los umbrales calibrados.',
8:'El criterio local reduce los sesgos respecto de imponer 1 mm en ALADIN. Para comparar local y global en las mismas 3 526 celdas, el sesgo firmado de la duración media pasa de −8,230 a −6,522 días y su error absoluto medio de 14,890 a 6,673 días; en t99, el sesgo pasa de −46,260 a −34,096 y el error absoluto medio de 71,473 a 48,825 días. El sesgo firmado permite compensaciones entre celdas de signo contrario; el error absoluto no. Estos valores sobre máscara común difieren de la Tabla S1, que presenta la máscara propia de cada criterio.',
9:'El mínimo de rachas se evaluó como sensibilidad después del análisis con 30. Con 50, 100 y 200 rachas, el sesgo medio de t99 pasa de −35,7 a −31,5, −20,3 y −10,5 días, y el de la duración media de −6,7 a −5,8, −4,4 y −2,8 días. Cambia también la cobertura espacial: se excluyen preferentemente celdas del norte. La reducción del sesgo no representa una corrección del modelo sobre un dominio idéntico.',
10:'Tabla S1b. Sensibilidad al mínimo de rachas por celda con calibración local. Porcentajes de celdas donde ALADIN es significativamente más corto; cada fila usa su propia máscara.',
11:'S2. Distribuciones históricas por ventana con distintos umbrales',
12:'Tabla S2. Duraciones históricas agrupadas por ventana, 1980–2014. CR2MET a 1 mm y ALADIN con 1 mm fijo, umbral global y campo local actualizado. n es el número de rachas completas; duraciones en días.',
17:'Figura S1. Sensibilidad histórica al umbral en Loa, Maipo, Maule y Biobío. CR2MET a 1 mm y ALADIN con 1 mm, 5,285 mm global y R₀(x,y) local. Bins comunes dentro de cada ventana; la integral es el número de rachas. El criterio local corresponde al campo completo, incluida la ampliación de búsqueda sobre 15 mm.',
19:'La respuesta varía según ventana y métrica. El campo local aproxima especialmente Maipo: media de 10,434 frente a 10,960 días y p99 de 121 frente a 119. En Loa, Maule y Biobío el p99 sigue siendo menor que el de CR2MET. Igualar la frecuencia húmeda anual no obliga a que coincidan las duraciones ni el número de rachas.',
23:'S3. Campos locales de umbral',
26:'Figura S2. Campos históricos de umbral de ALADIN calibrados por celda contra CR2MET a 1 mm y a 0,1 mm. Ambos cubren las 3 587 celdas. Cada campo iguala su propia frecuencia húmeda de referencia y se conserva fijo en el futuro.',
28:'Para la referencia de 1 mm, R₀ tiene mediana de 6,489 mm día⁻¹, p10 de 2,285 y p90 de 11,602, con rango 0,039–25,026. Se amplía el intervalo inicial de 0–15 mm en 82 celdas. Para 0,1 mm, la mediana es 4,241, p10 de 1,135 y p90 de 8,203, con rango 0,010–20,006; cinco celdas requieren ampliar la búsqueda. No quedan celdas sin calibrar en ninguno de los campos. El campo local anterior, limitado a 15 mm, no es la versión principal actual.',
30:'Los RR se recalculan imponiendo 1 mm en ALADIN histórico y futuro, sin calibrar su frecuencia. Se mantienen el inicio marzo–noviembre, las regiones y el bootstrap. Es una sensibilidad al criterio de clasificación, no una demostración de independencia respecto del umbral.',
32:'Tabla S3. RR(20) con campo local principal y con 1 mm fijo en ALADIN, SSP5-8.5. Conteos histórico → futuro e IC95. El criterio fijo conserva sesgos de frecuencia húmeda.',
37:'Figura S3. Curvas RR(D) con 1 mm fijo en ALADIN. IC95 puntuales por años; símbolos huecos si existen réplicas con N_hist* = 0. RR no definido cuando el conteo histórico observado es cero.',
39:'Con 1 mm fijo, Coquimbo y La Araucanía presentan aumentos significativos a 20 días. O’Higgins incluye 1. Los Lagos tiene solo 107 eventos históricos y 27 futuros, RR = 0,25, con un intervalo muy amplio que incluye 1: el valor es calculable, pero no respalda una disminución. La diferencia con el campo local confirma que la definición importa, especialmente donde los eventos largos son escasos.',
42:'Se contrasta la interpolación lineal a centros con el promedio de las celdas CR2MET cuyos centros caen dentro del polígono de cada celda ALADIN. Este segundo método no es un regrillado conservativo ponderado por solapamiento de áreas. Se recalibra cada campo local contra la referencia correspondiente. Con la referencia de 1 mm, el sesgo medio local cambia de −6,727 a −6,608 días y el de t99 de −35,662 a −34,234, sobre 3 549 y 3 552 celdas válidas, respectivamente. La comparación presenta máscaras propias; el signo medio se conserva.',
44:'Tabla S4. Interpolación lineal frente al promedio de centros dentro de celda. Sesgos espaciales en días y cobertura propia de cada combinación, CR2MET a 1 mm. iii es principal; global conserva el criterio i original.',
48:'Se repiten los RR con SSP3-7.0 y el mismo campo local histórico contra CR2MET a 1 mm. A 20 días, los cuatro cocientes son mayores que 1, pero solo La Araucanía queda marginalmente sobre 1: IC95 [1,0037; 2,5920]. No se interpreta esa cercanía al límite como robustez independiente del escenario. Un miembro por escenario no permite atribuir sus diferencias exclusivamente al forzamiento.',
49:'Tabla S5. RR(20) con calibración local, referencia de 1 mm y ambos escenarios. IC95 con cuatro decimales para evitar ocultar el límite cercano a 1 en La Araucanía bajo SSP3-7.0.',
53:'Figura S4. RR(D) bajo SSP3-7.0 con el campo local principal fijo, inicio marzo–noviembre. Bandas de IC95 puntuales; el denominador cero se trata como en la Figura S3.',
55:'Se compara la frecuencia seca de marzo–noviembre y ρ = P(seco hoy | seco ayer) − P(seco hoy | húmedo ayer). Se estima por celda y mes con suavizado de transiciones (+0,5 en numerador y +1 en denominador) y se promedia con igual peso sobre celdas y meses. Las diferencias usan 1 000 remuestreos independientes de años por período. Con la referencia principal de 1 mm, la frecuencia seca aumenta de forma detectable en las cuatro regiones y los IC95 de Δρ incluyen cero en todas. Ausencia de detección no equivale a igualdad; Los Lagos queda próximo al límite. La sensibilidad de 0,1 mm modifica este último resultado (S8).',
56:'Tabla S6. Frecuencia seca y persistencia de primer orden con calibración local contra 1 mm, histórico → futuro SSP5-8.5. Δ de frecuencia en puntos porcentuales e IC95; ρ e intervalos con cinco decimales en las diferencias.'
}
for i,t in su.items():setp(p[i],t)
s1=rows(A/'paper_figuras/tablaS1_sensibilidad_umbral_chile.csv')
table_fill(s.tables[0],[[v if j<2 else (integer(v) if j==8 else fmt(v)) for j,v in enumerate(r.values())] for r in s1])
mins=rows(A/'resultados_tesis/sensibilidad_minimo_rachas.csv')
table_fill(s.tables[1],[[r['minimo_rachas'],integer(r['celdas']),fmt(r['delta_media_promedio']),fmt(r['delta_t99_promedio']),fmt(r['delta_t99_mediana']),fmt(r['pct_media_ALADIN_menor_sig']),fmt(r['pct_t99_ALADIN_menor_sig'])] for r in mins])
hist=rows(A/'_pregunta9_calibracion_outputs/resumen_sensibilidad_umbral_4cuencas.csv')
print('HIST COLS',hist[0])
table_fill(s.tables[2],[[r['cuenca'],r['serie'],integer(r['n_spells']), *[fmt(r[k]) for k in ['media_dias','mediana_dias','p90_dias','p99_dias']],integer(r['max_dias'])] for r in hist])
s.tables[2].rows[0].cells[0].text='Ventana'
fixed=json.loads((A/'outputs_pregunta5_aladin_1mm/results.json').read_text())['rr_table']
table_fill(s.tables[3],[[r['Region'],f"{integer(r['N_hist'])} → {integer(r['N_fut'])}",f"{fmt(r['RR observado'],4)} [{fmt(r['IC95 inferior'],4)}; {fmt(r['IC95 superior'],4)}]",f"{integer(f['N_hist'])} → {integer(f['N_fut'])}",f"{fmt(f['RR observado'],4)} [{fmt(f['IC95 inferior'],4)}; {fmt(f['IC95 superior'],4)}]"] for r,f in zip(rr,fixed)])
setp(p[39],p[39].text+' En Los Lagos se descartan 42 de 1 000 réplicas a 20 días con el umbral fijo; su IC95 es condicional a las 958 restantes.')
for c,v in zip(s.tables[3].rows[0].cells,['Región','Local: N hist → fut','Local: RR [IC95]','Fijo 1 mm: N hist → fut','Fijo 1 mm: RR [IC95]']):c.text=v
rg=rows(A/'_test_regrillado/mapas_lineal_vs_area.csv')
for c,v in zip(s.tables[4].rows[0].cells,['Criterio / métrica','Lineal: Δ (celdas)','Promedio de centros: Δ (celdas)']):c.text=v
table_fill(s.tables[4],[[f"{k} / {m}",f"{fmt(next(r for r in rg if r['cr2met']=='lineal' and r['criterio']==k and r['metrica']==m)['Delta'],3)} ({next(r for r in rg if r['cr2met']=='lineal' and r['criterio']==k and r['metrica']==m)['pix']})",f"{fmt(next(r for r in rg if r['cr2met']=='area' and r['criterio']==k and r['metrica']==m)['Delta'],3)} ({next(r for r in rg if r['cr2met']=='area' and r['criterio']==k and r['metrica']==m)['pix']})"] for k in ['global','ii','iii'] for m in ['media','t99']])
ssp=json.loads((A/'outputs_pregunta5_aladin_ssp370/results.json').read_text())['rr_table']
table_fill(s.tables[5],[[r['Region'],integer(r['N_hist']),integer(r['N_fut']),f"{fmt(r['RR observado'],4)} [{fmt(r['IC95 inferior'],4)}; {fmt(r['IC95 superior'],4)}]",integer(f['N_fut']),f"{fmt(f['RR observado'],4)} [{fmt(f['IC95 inferior'],4)}; {fmt(f['IC95 superior'],4)}]"] for r,f in zip(rr,ssp)])
pers=rows(A/'resultados_tesis/contraste_persistencia_regiones.csv')
def persdata(data):return [[r['region'],f"{fmt(float(r['frec_seca_hist'])*100)} → {fmt(float(r['frec_seca_fut'])*100)} %; {fmt(float(r['frec_seca_delta'])*100,2)} [{fmt(float(r['frec_seca_delta_IC95_inf'])*100,2)}; {fmt(float(r['frec_seca_delta_IC95_sup'])*100,2)}]",f"{fmt(r['rho_hist'],3)} → {fmt(r['rho_fut'],3)}",f"{fmt(r['rho_delta'],5)} [{fmt(r['rho_delta_IC95_inf'],5)}; {fmt(r['rho_delta_IC95_sup'],5)}]",'Sí' if r['rho_IC_incluye_0']=='True' else 'No'] for r in data]
table_fill(s.tables[6],persdata(pers))
supppaths=[copyfig(A/'paper_figuras/figS1_sensibilidad_umbral_cuencas.png','Figura_S1.png'),copyfig(COMP/'comparacion_mapas_umbrales.png','Figura_S2.png'),copyfig(A/'paper_figuras/figS_rr_curves_aladin_1mm.png','Figura_S3.png'),copyfig(A/'paper_figuras/figS_rr_curves_aladin_ssp370.png','Figura_S4.png')]
replace_images(s,supppaths)
s.add_heading('S8. Sensibilidad a la referencia de 0,1 mm',level=2)
s.add_paragraph('La referencia principal de 1 mm sigue la convención ETCCDI de día húmedo (Zhang et al., 2011) y facilita la comparación con Latoja et al. (2025). No se elige por maximizar la significancia. Una referencia de 0,1 mm identifica más interrupciones por lluvia débil y tiene usos válidos, por ejemplo en Breinl et al. (2020). Se repite la cadena completa con CR2MET a 0,1 mm y una calibración histórica propia por celda; no se reutiliza el campo de 1 mm ni se recalibra en el futuro. Se conserva el resto de las reglas del artículo.')
s.add_paragraph('Tabla S7. Comparación histórica sobre las mismas 3 549 celdas para ambas referencias. Cada referencia cambia también las duraciones observadas; una reducción del error en días no demuestra que una definición sea físicamente superior. Δ = ALADIN − CR2MET; MAE es el promedio de |Δ| por celda.')
sp=rows(COMP/'04_resumen_espacial.csv');sp=[r for r in sp if r['muestra']=='mismas celdas']
new_table(s,['Métrica','Referencia','CR2MET (d)','ALADIN (d)','Δ (d)','MAE (d)','Error relativo absoluto (%)'],[[('Media' if r['metrica']=='mean' else 't99'),ref,fmt(r['CR2MET_'+suffix],3),fmt(r['ALADIN_'+suffix],3),fmt(r['sesgo_'+suffix],3),fmt(r['MAE_'+suffix],3),fmt(r['error_rel_absoluto_'+suffix+'_pct'],3)] for r in sp for ref,suffix in [('1 mm','1mm'),('0,1 mm','0p1mm')]])
s.add_paragraph('La máscara propia de la referencia de 0,1 mm contiene 3 583 celdas con al menos 30 rachas, frente a 3 549 con 1 mm. La Tabla S7 compara exclusivamente las 3 549 comunes; sus promedios no deben sustituirse por los de las máscaras propias.')
s.add_paragraph('Tabla S8. RR(20) con calibración por celda para ambas referencias y escenarios. IC95 puntuales; descartes entre 1 000 réplicas por ausencia de eventos históricos. Los intervalos con descartes están condicionados a las réplicas restantes.')
cmp=[r for r in rows(COMP/'09_RR_D20.csv') if r['escenario'] in ['SSP5-8.5','SSP3-7.0']]
new_table(s,['Escenario','Región','Referencia','N hist → fut','RR [IC95]','Descartes'],[[r['escenario'],r['region'],ref,f"{integer(r['N_hist_'+suffix])} → {integer(r['N_fut_'+suffix])}",f"{fmt(r['RR_'+suffix],4)} [{fmt(r['IC_inf_'+suffix],4)}; {fmt(r['IC_sup_'+suffix],4)}]",r['descartadas_'+suffix]] for r in cmp for ref,suffix in [('1 mm','1mm'),('0,1 mm','0p1mm')]])
s.add_paragraph('Con 0,1 mm y SSP5-8.5 son significativos Coquimbo y La Araucanía. En Coquimbo, los conteos históricos bajan de 26 613 a 24 623 y los futuros de 28 788 a 27 647; el denominador disminuye proporcionalmente más y el RR pasa de 1,0817 a 1,1228. Es una descripción aritmética del cambio de definición, no una atribución física. Sus intervalos se solapan: la diferencia entre decisiones de significancia no prueba una diferencia significativa entre ambos RR. Con SSP3-7.0, ninguno de los cuatro RR(20) es significativo bajo la referencia de 0,1 mm. Los Lagos cambia incluso el signo del cociente puntual en SSP5-8.5, pero ambos intervalos incluyen 1 y son muy amplios.')
add_fig(s,copyfig(COMP/'comparacion_RR.png','Figura_S5.png'),'Figura S5. Comparación de curvas RR(D) para referencias de 1 y 0,1 mm con calibración por celda. Los intervalos se interpretan según la disponibilidad de denominadores históricos; no se concluyen diferencias entre criterios a partir del solapamiento visual.')
s.add_paragraph('Tabla S9. Frecuencia seca y persistencia con referencia de 0,1 mm, SSP5-8.5. En Los Lagos, Δρ = 0,02800 [0,00387; 0,05209] excluye cero, a diferencia de 1 mm; por tanto, la ausencia de cambio detectable de primer orden no es robusta a ambas definiciones en esa región.')
pers01=rows(B/'resultados_tesis/contraste_persistencia_regiones.csv')
new_table(s,['Región','Frecuencia seca hist → fut; Δ [IC95], pp','ρ hist → fut','Δρ [IC95]','Incluye 0'],persdata(pers01))
s.add_heading('S9. Referencia de Markov y tratamiento de los bordes',level=2)
s.add_paragraph('Se usa una cadena de Markov mensual de primer orden con 40 simulaciones por celda (semilla 7). La referencia de solo frecuencia modifica la frecuencia seca futura manteniendo el parámetro de persistencia histórico. Se compara su RR puntual con el IC95 del conteo observado, sin interpretar la compatibilidad como prueba causal. La dispersión entre simulaciones representa ruido Monte Carlo y no incertidumbre de los parámetros. Las réplicas bootstrap de persistencia usan semilla 11.')
mk=rows(A/'resultados_tesis/descomposicion_markov_rr_regiones.csv');mk01=rows(B/'resultados_tesis/descomposicion_markov_rr_regiones.csv')
s.add_paragraph('Con 1 mm, la referencia de solo frecuencia está dentro del IC95 observado en 30 de 30 combinaciones región–duración con N_hist ≥ 100; con 0,1 mm, en 29 de 29. Esta selección evita interpretar cocientes de eventos extremadamente escasos, pero no valida por sí sola la distribución histórica de la cadena.')
s.add_paragraph('Tabla S10. Diagnóstico Markov a 20 días. Cobertura = número histórico simulado / número histórico observado. Se reporta para evaluar si la compatibilidad futura descansa en una representación histórica suficiente.')
new_table(s,['Región','Referencia','N histórico observado','Cobertura histórica','RR observado','RR solo frecuencia'],[[r['region'],ref,integer(r['N_hist_obs']),fmt(r['cobertura_markov_hist'],3),fmt(r['RR_obs'],3),fmt(r['RR_solo_frec'],3)] for ref,data in [('1 mm',mk),('0,1 mm',mk01)] for r in data if float(r['D'])==20])
s.add_paragraph('Tabla S11. Rachas censuradas en las ventanas con referencia de 1 mm. El porcentaje usa completas + censuradas como denominador. La última columna cuenta duraciones observadas truncadas que superan el p99 de las completas; sus duraciones totales no se conocen.')
cen=rows(A/'_test_regrillado/rachas_censuradas_por_cuenca.csv')
new_table(s,['Serie','Ventana','Completas','Censuradas','Porcentaje','Truncadas > p99'],[[r['serie'],r['cuenca'],integer(r['rachas_completas']),r['censuradas'],fmt(r['pct_censuradas'],3),r['censuradas_sobre_p99']] for r in cen])
s.add_paragraph('Reproducibilidad de la sensibilidad de 0,1 mm. La precipitación CR2MET interpolada se conserva con precisión float64 antes de clasificar. La cuantización anticipada a float32 puede cambiar días extremadamente próximos a 0,1 mm. El análisis auxiliar de regrillado y censura se regeneró con esa precisión y se contrastó con la cadena principal; el campo calibrado se almacena en float32 sin redondearlo a tres decimales. La versión limitada a 15 mm y las salidas anteriores se conservan como antecedentes, pero no se mezclan con las tablas de esta versión.')
s.add_heading('S10. Comparación de calibración global y local',level=2)
s.add_paragraph('Tabla S12. RR(20) bajo SSP5-8.5 para las cuatro combinaciones de referencia y calibración. Los umbrales globales auxiliares son 5,285 mm para la referencia de 1 mm y 3,668 mm para 0,1 mm. Cada método se aplica con su propio umbral histórico fijo en el futuro. Se presentan todos los resultados, sin seleccionar la definición por su significancia.')
four=[r for r in rows(COMP/'18_comparacion_cuatro_versiones_RR.csv') if float(r['D'])==20 and r['escenario']=='SSP5-8.5']
new_table(s,['Referencia','Calibración','Región','N hist → fut','RR [IC95]','Descartes'],[[r['referencia'],r['calibracion'],r['region'],f"{integer(r['N_hist'])} → {integer(r['N_fut'])}",f"{fmt(r['RR'],4)} [{fmt(r['IC_inf'],4)}; {fmt(r['IC_sup'],4)}]",r['descartadas']] for r in four])
s.add_paragraph('El cambio de significancia de Coquimbo al pasar del umbral global al local con referencia de 1 mm muestra que la comparación entre regiones depende de cómo se controla la frecuencia húmeda. El objetivo del criterio local es igualar esa frecuencia en cada celda; no conservar un resultado inferencial de la versión global. Del mismo modo, el resultado con 0,1 mm se interpreta como sensibilidad a otra definición de día húmedo.')
s.add_heading('Referencias del material suplementario',level=2)
s.add_paragraph('Breinl, K., Di Baldassarre, G., Mazzoleni, M., Lun, D., y Vico, G. (2020). Extreme dry and wet spells face changes in their duration and timing. Environmental Research Letters, 15(7), 074040. https://doi.org/10.1088/1748-9326/ab7d05')
for idx in [121,138]:s.add_paragraph(Document(paper_src).paragraphs[idx].text)
for para in s.paragraphs:
 if not para.text:continue
 is_title=para.text=='Material suplementario'
 is_heading=bool(re.match(r'^S\d+\.',para.text)) or para.text=='Referencias del material suplementario'
 para.paragraph_format.line_spacing=1.5
 para.paragraph_format.space_after=Pt(6)
 if is_heading:
  para.style=s.styles['Heading 2'];para.paragraph_format.keep_with_next=True
 for run in para.runs:
  run.font.name='Calibri';run.font.size=Pt(16 if is_title else 12)
  run.font.italic=False;run.font.bold=is_heading
  run.font.color.rgb=RGBColor(0,0,0)
 if para.text.startswith(('Tabla S','Figura S')):
  para.paragraph_format.keep_with_next=para.text.startswith('Tabla S')
  for run in para.runs:run.font.size=Pt(11)
clean(s);s.save(OUT/'Material_suplementario_umbral_local.docx')
(OUT/'figuras/INDICE_FIGURAS.json').write_text(json.dumps(FIG,ensure_ascii=False,indent=2))
print('CREATED',len(d.paragraphs),len(s.paragraphs),len(s.tables))
