from pathlib import Path
from zipfile import ZipFile
from lxml import etree
from docx import Document
import json,hashlib,csv,re
BASE=Path(__file__).resolve().parent
ROOT=BASE.parent
def readcsv(p):
 with open(p) as f:return list(csv.DictReader(f))
def f(x,n=1):return f'{float(x):.{n}f}'.replace('.',',')
def integer(x):return f'{int(float(x)):,}'.replace(',',' ')
checks=[]
def check(label,result):
 checks.append({'control':label,'ok':bool(result)})
 if not result:raise AssertionError(label)
d=Document(BASE/'Paper_umbral_local_1mm.docx');s=Document(BASE/'Material_suplementario_umbral_local.docx')
A=ROOT/'temporal_umbral_por_celda_2026-10-24/ejecucion'
check('Cuatro tablas y seis figuras en el manuscrito',len(d.tables)==4 and len(d.inline_shapes)==6)
check('Trece tablas y cinco figuras en el suplemento',len(s.tables)==13 and len(s.inline_shapes)==5)
for row,src in zip(d.tables[2].rows[1:],readcsv(A/'resultados_tesis/tabla3_estadisticos_ventanas.csv')):
 expected=[src['ventana'],src['producto'],integer(src['n_spells']),f(src['media']),f(src['mediana']),f(src['p90']),f(src['p99']),integer(src['max']),f(src['frac_ge20_pct'])]
 check('Tabla 3 '+src['ventana']+' '+src['producto'],[c.text for c in row.cells]==expected)
rr=json.loads((A/'outputs_pregunta5_aladin/results.json').read_text())['rr_table']
for row,src in zip(d.tables[3].rows[1:],rr):
 vals=[c.text for c in row.cells]
 check('Tabla 4 '+src['Region'],vals==[src['Region'],integer(src['N_hist']),integer(src['N_fut']),f(src['RR observado'],4),f"[{f(src['IC95 inferior'],4)}; {f(src['IC95 superior'],4)}]",'Sí' if src['Sig aumento'] else 'No'])
check('Solo Araucanía significativa en el principal',[r['Region'] for r in rr if r['Sig aumento']]==['La Araucanía'])
for name,doc in [('paper',d),('suplemento',s)]:
 text='\n'.join(p.text for p in doc.paragraphs)
 check(name+' sin notas pendientes ni volcado xarray',not any(x in text for x in ['[VERIFICAR','[PENDIENTE','xarray.DataArray']))
 check(name+' conserva geometría de página',len(doc.sections)==1 and doc.sections[0].page_width.inches==8.5 and doc.sections[0].page_height.inches==11)
 check(name+' incluye sensibilidad de Los Lagos','0,1 mm' in text and 'Los Lagos' in text and ('0,02800' in text or 'aumento detectable de ρ' in text))
 check(name+' celdas comunes identificadas','3 549' in text and '3 526' in text)
 for t in doc.tables:
  for r in t.rows:check(name+' fila protegida '+str(len(checks)),r._tr.trPr.find('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}cantSplit') is not None)
for name in ['paper','suplemento']:
 info=json.loads((BASE/'qa'/f'{name}_inventario.json').read_text())
 src=Path(info['source'])
 check(name+' original inalterado',hashlib.sha256(src.read_bytes()).hexdigest()==info['sha256'])
 check(name+' copia fuente idéntica',hashlib.sha256((BASE/'fuentes'/src.name).read_bytes()).hexdigest()==info['sha256'])
source=ZipFile(BASE/'fuentes'/Path(json.loads((BASE/'qa/paper_inventario.json').read_text())['source']).name)
final=ZipFile(BASE/'Paper_umbral_local_1mm.docx')
check('Comentarios archivados y eliminados solo de la entrega','word/comments.xml' in source.namelist() and 'word/comments.xml' not in final.namelist())
for info in json.loads((BASE/'figuras/INDICE_FIGURAS.json').read_text()):
 check('Integridad figura '+info['archivo'],hashlib.sha256((BASE/'figuras'/info['archivo']).read_bytes()).hexdigest()==info['sha256'])
# Guardar comparación de partes; los cambios autorizados afectan contenido/estilos/medios y eliminación de comentarios.
changes={}
for label,file,inventory in [('paper','Paper_umbral_local_1mm.docx','paper'),('suplemento','Material_suplementario_umbral_local.docx','suplemento')]:
 original=json.loads((BASE/'qa'/f'{inventory}_inventario.json').read_text())['parts'];z=ZipFile(BASE/file)
 now={n:hashlib.sha256(z.read(n)).hexdigest() for n in z.namelist()}
 changes[label]={'cambiadas':[n for n in original if n in now and original[n]!=now[n]],'eliminadas':[n for n in original if n not in now],'nuevas':[n for n in now if n not in original]}
(BASE/'qa/controles_documentos.json').write_text(json.dumps({'total':len(checks),'correctos':sum(c['ok'] for c in checks),'controles':checks,'comparacion_paquetes':changes},ensure_ascii=False,indent=2))
print(json.dumps({'total':len(checks),'correctos':sum(c['ok'] for c in checks),'partes':changes},ensure_ascii=False))
